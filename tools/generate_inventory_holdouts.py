"""Generate synthetic unfamiliar-layout regression inputs, not provider certification."""
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "api/tests/Advertified.Commercial.Api.Tests/Fixtures/InventoryHoldouts"
STAMP = datetime(2026, 9, 12)
HEADERS = ["Asset key", "Placement description", "Quoted amount", "Money unit"]
MEANINGS = ["product_code", "name", "rate", "currency"]


def product_rows(tag):
    return [
        [f"{tag}-731", "Synthetic east panel", "1847.35", "ZAR"],
        [f"{tag}-946", "Synthetic west panel", "2936.42", "ZAR"],
    ]


def deterministic_zip(path):
    with ZipFile(path) as archive:
        entries = [(item.filename, archive.read(item)) for item in archive.infolist()]
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in sorted(entries):
            item = ZipInfo(name, (2026, 9, 12, 0, 0, 0))
            item.compress_type = ZIP_DEFLATED
            archive.writestr(item, content)


def table(rows, first=2, span=1, axis="ROW", meanings=None, header=1):
    return dict(rows=rows, first=first, last=len(rows) if axis == "ROW" else len(rows[0]),
                span=span, axis=axis, header=header,
                meanings=meanings or MEANINGS)


def write_sheet(path, tables, merge=False):
    book = Workbook()
    book.remove(book.active)
    book.properties.created = STAMP
    book.properties.modified = STAMP
    for index, spec in enumerate(tables, 1):
        sheet = book.create_sheet(f"Unseen section {index}")
        for row in spec["rows"]:
            sheet.append(row)
        if merge:
            sheet.merge_cells("A1:D1")
    book.save(path)
    deterministic_zip(path)


def tabular_cases():
    result = []
    for key in ["flat", "multi_headers", "merged_header", "transposed",
                "multi_section", "package_schedule", "missing_supplier", "rate_variants"]:
        rows = [HEADERS, *product_rows(key)]
        spec = table(rows)
        if key in ("multi_headers", "merged_header"):
            spec = table([["Synthetic inventory catalogue"], HEADERS, *product_rows(key)],
                         first=3, header=2)
        if key == "transposed":
            spec = table(list(map(list, zip(*rows))), axis="COLUMN")
        if key == "package_schedule":
            spec = table([HEADERS, rows[1], ["", "Synthetic Tuesday schedule", "", ""],
                          rows[2], ["", "Synthetic Friday schedule", "", ""]], span=2)
        if key == "rate_variants":
            spec = table([HEADERS + ["Alternative duration amount"],
                          rows[1] + ["8271.19"], rows[2] + ["9182.63"]],
                         meanings=MEANINGS + ["rate"])
        tables = [spec]
        if key == "multi_section":
            tables.append(table([HEADERS, *product_rows("other")]))
        path = OUT / f"xlsx_{key}.xlsx"
        write_sheet(path, tables, merge=key == "merged_header")
        result.append(dict(file=path.name, format="XLSX", tables=tables))
    for key in ["quoted_commas", "blank_cells", "extra_columns", "unusual_headers"]:
        rows = [HEADERS.copy(), *product_rows(key)]
        meanings = MEANINGS.copy()
        if key == "quoted_commas":
            rows[1][1] = 'Synthetic panel, east "wing"'
        if key == "blank_cells":
            rows[1][2] = ""
            rows[2][1] = ""
        if key == "extra_columns":
            rows = [rows[0] + ["Survey note"], rows[1] + ["Unverified access"],
                    rows[2] + ["No field inspection"]]
            meanings += [None]
        if key == "unusual_headers":
            rows[0] = ["Reference / item", "Exposure surface", "Offer in money", "ISO tender"]
        path = OUT / f"csv_{key}.csv"
        with path.open("w", newline="", encoding="utf-8") as stream:
            csv.writer(stream, lineterminator="\n").writerows(rows)
        result.append(dict(file=path.name, format="CSV", tables=[table(rows, meanings=meanings)]))
    return result


def write_slides(path, pages):
    deck = Presentation()
    deck.core_properties.created = STAMP
    deck.core_properties.modified = STAMP
    for lines in pages:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for index, line in enumerate(lines):
            box = slide.shapes.add_textbox(Inches(.5), Inches(.5 + index * .6),
                                          Inches(8), Inches(.5))
            box.text_frame.text = line
    deck.save(path)
    deterministic_zip(path)


def write_pdf(path, pages, columns=False):
    canvas = Canvas(str(path), invariant=1, pageCompression=1)
    for lines in pages:
        for index, line in enumerate(lines):
            x = 45 + (275 if columns and index % 2 else 0)
            y = 780 - (index // 2 if columns else index) * 32
            canvas.drawString(x, y, line)
        canvas.showPage()
    canvas.save()


def narrative_cases():
    result = []
    cases = [
        ("PPTX", "repeated_products", [["Synthetic panel H-591", "Quoted ZAR 2719.38"],
                                      ["Synthetic panel J-683", "Quoted ZAR 3826.47"]]),
        ("PPTX", "schedule", [["Synthetic package K-752", "Tuesday 09:00 slot",
                              "Friday 18:00 slot", "Bundle ZAR 9137.24"]]),
        ("PPTX", "narrative_values", [["Synthetic panel N-815", "Inspection not supplied",
                                      "Indicative ZAR 4271.69", "Availability not supplied"]]),
        ("PDF", "text_rate_card", [["Synthetic panel P-937", "Quoted ZAR 5731.28"]]),
        ("PDF", "brochure", [["Synthetic brochure B-463", "Supplier not supplied"],
                             ["Synthetic panel C-529", "Indicative ZAR 6217.39"]]),
        ("PDF", "mixed_narrative_table", [["Synthetic catalogue M-647", "Rates exclude installation",
                                          "M-647 | Wall | ZAR 1738.42",
                                          "M-748 | Screen | ZAR 2849.53"]]),
        ("PDF", "difficult_layout", [["Synthetic left L-891", "Synthetic right R-927",
                                     "Left ZAR 3726.41", "Right ZAR 4837.52",
                                     "Dates not supplied", "Availability not supplied"]]),
    ]
    for kind, key, pages in cases:
        path = OUT / f"{kind.lower()}_{key}.{kind.lower()}"
        if kind == "PPTX":
            write_slides(path, pages)
        else:
            write_pdf(path, pages, columns=key == "difficult_layout")
        result.append(dict(file=path.name, format=kind, pages=pages))
    return result


def main():
    OUT.mkdir(parents=True, exist_ok=False)
    cases = tabular_cases() + narrative_cases()
    for case in cases:
        path = OUT / case["file"]
        case["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        case["bytes"] = path.stat().st_size
    manifest = dict(schema="advertified.synthetic-inventory-holdouts.v1",
                    generated_by="tools/generate_inventory_holdouts.py",
                    limitation="Deterministic mappings prove source fidelity, not model generalization.",
                    cases=cases)
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(dict(cases=len(cases), bytes=sum(c["bytes"] for c in cases))))


if __name__ == "__main__":
    main()
